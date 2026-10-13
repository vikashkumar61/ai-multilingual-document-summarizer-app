"""Extract text and basic metadata from PDF, DOCX, and TXT files."""
from pathlib import Path

import pymupdf
from docx import Document

from preprocessing import clean_text


def extract_text(file_obj, return_metadata=False, filename="document.txt"):
    name = (filename or "document.txt").lower()
    raw = file_obj.read()
    if hasattr(file_obj, "seek"):
        file_obj.seek(0)

    pages = 0
    if name.endswith(".pdf"):
        doc = pymupdf.open(stream=raw, filetype="pdf")
        try:
            pages = len(doc)
            text = "\n".join(page.get_text("text") for page in doc)
        finally:
            doc.close()
    elif name.endswith(".docx"):
        from io import BytesIO
        document = Document(BytesIO(raw))
        text = "\n".join(p.text for p in document.paragraphs)
        # Include simple table content too.
        table_lines = []
        for table in document.tables:
            for row in table.rows:
                table_lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        if table_lines:
            text += "\n" + "\n".join(table_lines)
    elif name.endswith(".txt"):
        text = raw.decode("utf-8-sig", errors="replace")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")

    text = clean_text(text)
    if not text:
        raise ValueError(
            "No readable text was found. If this is a scanned PDF, OCR is required."
        )
    metadata = {
        "pages": pages,
        "characters": len(text),
    }
    return (text, metadata) if return_metadata else text
